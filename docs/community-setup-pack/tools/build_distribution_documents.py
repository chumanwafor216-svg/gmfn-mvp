from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

DOC_DATE = "2026-09-14"
PACK_NAME = "GSN Community Setup Pack"
AUTHENTICITY_LINE = (
    f"GSN Original Community Setup Material | Source-controlled in gmfn_mvp/docs | Document date: {DOC_DATE}"
)
WATERMARK_TEXT = f"GSN ORIGINAL | {DOC_DATE}"

REPO_ROOT = Path(__file__).resolve().parents[3]
GUIDE_DIR = REPO_ROOT / "docs" / "gsn-user-guide"
REAL_LIFE_DIR = GUIDE_DIR / "real-life-capability-bank"

TITLE_MAP = {
    GUIDE_DIR / "GSN_SELF_SERVICE_USER_GUIDE.md": (
        f"{PACK_NAME}: What to Set Up, What It Does, and How to Use It ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_MASTER_CAPABILITY_BANK.md": (
        f"{PACK_NAME}: GSN in Real Life Master Capability Bank ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_AUDIENCE_PICK_LISTS.md": (
        f"{PACK_NAME}: Audience Pick Lists for Community Setup ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_CHURCH_FAITH_GROUP_COPY.md": (
        f"{PACK_NAME}: Church and Faith Group Real-Life Copy ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_SCHOOL_YOUTH_ORGANISATION_COPY.md": (
        f"{PACK_NAME}: School and Youth Organisation Real-Life Copy ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_NGO_CHARITY_COMMUNITY_ORGANISATION_COPY.md": (
        f"{PACK_NAME}: NGO, Charity and Community Organisation Real-Life Copy ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_COOPERATIVE_MARKET_BUSINESS_NETWORK_COPY.md": (
        f"{PACK_NAME}: Cooperative, Market and Business Network Real-Life Copy ({DOC_DATE})"
    ),
    REAL_LIFE_DIR / "GSN_IN_REAL_LIFE_FAMILY_DIASPORA_ASSOCIATION_COPY.md": (
        f"{PACK_NAME}: Family, Diaspora and Association Real-Life Copy ({DOC_DATE})"
    ),
}


def ensure_markdown_identity(path: Path, title: str) -> None:
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    if lines and lines[0].startswith("# "):
        lines[0] = f"# {title}"
    else:
        lines.insert(0, f"# {title}")

    metadata = [
        f"Document date: {DOC_DATE}",
        f"Pack: {PACK_NAME}",
        f"Authenticity mark: {AUTHENTICITY_LINE}",
    ]
    body = "\n".join(lines).strip() + "\n"
    existing_lines = set(body.splitlines())
    for item in reversed(metadata):
        if item not in existing_lines:
            body = body.replace("\n\n", f"\n\n{item}\n", 1)
            existing_lines.add(item)
    path.write_text(body, encoding="utf-8")


def markdown_to_docx(md_path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        header = section.header.paragraphs[0]
        header.text = f"{PACK_NAME} | {DOC_DATE}"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        footer.text = f"{AUTHENTICITY_LINE} | Page "
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        footer._p.append(page_field)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)

    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:])
            run.font.name = "Aptos Display"
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(11, 37, 69)
            run.bold = True
        elif line.startswith("## "):
            paragraph = doc.add_heading(line[3:], level=1)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(11, 37, 69)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].replace("**", ""), style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", line), style="List Number")
        else:
            doc.add_paragraph(line.replace("**", ""))

    out = md_path.with_suffix(".docx")
    doc.save(out)
    add_docx_watermark(out)
    return out


def add_docx_watermark(docx_path: Path) -> None:
    try:
        from lxml import etree
    except Exception:
        return

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    v_ns = "urn:schemas-microsoft-com:vml"
    o_ns = "urn:schemas-microsoft-com:office:office"

    def tag(ns: str, name: str) -> str:
        return f"{{{ns}}}{name}"

    with zipfile.ZipFile(docx_path, "r") as zin:
        headers = [
            name for name in zin.namelist() if re.fullmatch(r"word/header\d+\.xml", name)
        ]
        if not headers:
            return
        header = sorted(headers, key=lambda name: int(re.findall(r"\d+", name)[0]))[0]
        root = etree.fromstring(zin.read(header))
        paragraph = etree.Element(tag(w_ns, "p"))
        run = etree.SubElement(paragraph, tag(w_ns, "r"))
        pict = etree.SubElement(run, tag(w_ns, "pict"))
        shape = etree.SubElement(
            pict,
            tag(v_ns, "shape"),
            {
                "id": "GSNCommunitySetupWatermark",
                tag(o_ns, "spid"): "_x0000_s1025",
                "type": "#_x0000_t136",
                "style": (
                    "position:absolute;margin-left:0;margin-top:0;"
                    "width:468pt;height:468pt;rotation:315;"
                    "z-index:-251654144;"
                    "mso-position-horizontal:center;"
                    "mso-position-vertical:center;mso-wrap-edited:f;"
                ),
                "fillcolor": "#C9CED6",
                "stroked": "f",
            },
        )
        etree.SubElement(shape, tag(v_ns, "fill"), {"opacity": "0.12"})
        etree.SubElement(
            shape,
            tag(v_ns, "textpath"),
            {"style": 'font-family:"Aptos";font-size:1pt', "string": WATERMARK_TEXT},
        )
        etree.SubElement(shape, tag(v_ns, "path"), {"textpathok": "t"})
        root.append(paragraph)
        overrides = {
            header: etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
        }
        parts = [(info, zin.read(info.filename)) for info in zin.infolist()]

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in parts:
            zout.writestr(info, overrides.get(info.filename, data))


def markdown_to_pdf(md_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    def esc(value: str) -> str:
        return (
            value.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("**", "")
        )

    out = md_path.with_suffix(".pdf")
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.7,
        leading=10.8,
        spaceAfter=5,
    )
    h1 = ParagraphStyle(
        "H1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0B2545"),
        spaceBefore=10,
        spaceAfter=6,
    )
    title = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=20,
        textColor=colors.HexColor("#0B2545"),
        spaceAfter=8,
    )

    story = []
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 3))
        elif line.startswith("# "):
            story.append(Paragraph(esc(line[2:]), title))
        elif line.startswith("## "):
            story.append(Paragraph(esc(line[3:]), h1))
        elif line.startswith("### "):
            story.append(Paragraph(esc(line[4:]), h1))
        elif line.startswith("- "):
            story.append(Paragraph("- " + esc(line[2:]), body))
        else:
            story.append(Paragraph(esc(line), body))

    def page_mark(canvas, doc_obj):
        canvas.saveState()
        canvas.setFillColor(colors.Color(0.08, 0.14, 0.22, alpha=0.08))
        canvas.setFont("Helvetica-Bold", 34)
        canvas.translate(letter[0] / 2, letter[1] / 2)
        canvas.rotate(38)
        canvas.drawCentredString(0, 0, WATERMARK_TEXT)
        canvas.restoreState()

        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.setFont("Helvetica", 7)
        canvas.drawCentredString(letter[0] / 2, 0.35 * inch, AUTHENTICITY_LINE)
        canvas.drawRightString(letter[0] - 0.75 * inch, 0.35 * inch, f"Page {doc_obj.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        str(out),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    ).build(story, onFirstPage=page_mark, onLaterPages=page_mark)
    return out


def main() -> None:
    rendered = []
    for path, title in TITLE_MAP.items():
        if not path.exists():
            continue
        ensure_markdown_identity(path, title)
        rendered.append(str(markdown_to_docx(path).relative_to(REPO_ROOT)))
        rendered.append(str(markdown_to_pdf(path).relative_to(REPO_ROOT)))

    print(json.dumps({"document_date": DOC_DATE, "rendered": rendered}, indent=2))


if __name__ == "__main__":
    main()
