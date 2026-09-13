import csv
import json
from pathlib import Path

ROOT = Path('docs/gsn-user-guide/real-life-capability-bank')
SOURCE = ROOT / 'capabilities.jsonl'
TITLE = 'GSN in Real Life'
SUBTITLE = 'Practical ways GSN helps communities organise, preserve value and create opportunity'
GUIDE_LINK = '../GSN_SELF_SERVICE_USER_GUIDE.md'

AUDIENCES = {
    'church': ('Church and Faith Group Copy', 'For pastors, church executives, departments, fellowships and faith-based community leaders.'),
    'school': ('School and Youth Organisation Copy', 'For schools, youth organisations, parent groups, alumni bodies and education programmes.'),
    'ngo': ('NGO, Charity and Community Organisation Copy', 'For charities, NGOs, programme teams, volunteers, trustees and funder-facing community work.'),
    'cooperative': ('Cooperative, Market and Business Network Copy', 'For cooperatives, traders, market associations, seller groups and local enterprise networks.'),
    'family': ('Family, Diaspora and Association Copy', 'For families, diaspora groups, town unions, clans, alumni circles and member associations.'),
}

COPY_NAME = {
    'church': 'GSN_IN_REAL_LIFE_CHURCH_FAITH_GROUP_COPY',
    'school': 'GSN_IN_REAL_LIFE_SCHOOL_YOUTH_ORGANISATION_COPY',
    'ngo': 'GSN_IN_REAL_LIFE_NGO_CHARITY_COMMUNITY_ORGANISATION_COPY',
    'cooperative': 'GSN_IN_REAL_LIFE_COOPERATIVE_MARKET_BUSINESS_NETWORK_COPY',
    'family': 'GSN_IN_REAL_LIFE_FAMILY_DIASPORA_ASSOCIATION_COPY',
}


def load_items():
    items = []
    with SOURCE.open('r', encoding='utf-8-sig') as fh:
        for line in fh:
            line = line.strip()
            if line:
                items.append(json.loads(line))
    return sorted(items, key=lambda x: x['id'])


def module_text(item, index_label=None):
    number = index_label if index_label is not None else item['id']
    lines = [f"## {number}. {item['title']}", '']
    lines += [f"**Current truth boundary:** {item['status']}", '']
    lines += [f"**The real-life problem:** {item['problem']}", '']
    lines += [f"**What GSN adds:** {item['adds']}", '']
    lines += [f"**Why it is important:** {item['why']}", '']
    lines += ["**Why the reader should care:**", '']
    lines += ["- **For management:** It gives leaders a cleaner process to supervise, delegate, prove and hand over instead of relying only on memory or scattered messages."]
    lines += ["- **For members:** It gives ordinary members a clearer way to see, respond, participate, sell, ask, verify or follow up without being trapped in long chat confusion."]
    lines += ["- **For the community body:** It turns activity into reusable memory, evidence and opportunity that can survive changing executives, phones and WhatsApp groups.", '']
    lines += [f"**The small detail that makes it different:** {item['details']}", '']
    lines += [f"**What not to overclaim:** {item['boundary']}", '']
    return '\n'.join(lines)


def write_md(path, title, intro, items, selected=False):
    lines = [f"# {title}", '', SUBTITLE, '', f"Linked setup/status authority: [{GUIDE_LINK}]({GUIDE_LINK})", '']
    lines += ["## How to Use This Document", '']
    lines += [intro, '']
    lines += ["This is a capability and messaging guide, not the final feature-status authority. Some items are live, some are partial, and some are capability directions that must be checked against the verified self-service guide before public claims or handover.", '']
    lines += ["Core principle: GSN presents organised evidence; the receiving person, organisation or community decides.", '']
    if selected:
        lines += ["Use this copy as an audience-specific conversation pack. It deliberately excludes many modules from the master bank so the receiver sees only what concerns them.", '']
    else:
        lines += ["Use this master bank like a menu. Do not send it all to everyone. Pick the modules that match the church, school, NGO, cooperative, market, family or diaspora conversation.", '']
    lines += ["## Capability Modules", '']
    for idx, item in enumerate(items, 1):
        lines.append(module_text(item, idx if selected else None))
    path.write_text('\n'.join(lines).rstrip() + '\n', encoding='utf-8')


def simple_docx(md_path):
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception as exc:
        return f'skipped docx: {exc}'
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(9.5)
    for raw in md_path.read_text(encoding='utf-8').splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:].replace('**',''), style='List Bullet')
        else:
            doc.add_paragraph(line.replace('**',''))
    out = md_path.with_suffix('.docx')
    doc.save(out)
    return str(out)


def simple_pdf(md_path):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except Exception as exc:
        return f'skipped pdf: {exc}'
    out = md_path.with_suffix('.pdf')
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=0.55*inch, rightMargin=0.55*inch, topMargin=0.55*inch, bottomMargin=0.55*inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='TinyBody', parent=styles['BodyText'], fontName='Helvetica', fontSize=8.2, leading=10.2, spaceAfter=4))
    styles.add(ParagraphStyle(name='SmallH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=15, leading=18, spaceAfter=7))
    styles.add(ParagraphStyle(name='SmallH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=13, spaceAfter=5))
    story = []
    for raw in md_path.read_text(encoding='utf-8').splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 3))
            continue
        clean = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('**', '')
        if clean.startswith('# '):
            story.append(Paragraph(clean[2:], styles['SmallH1']))
        elif clean.startswith('## '):
            story.append(Paragraph(clean[3:], styles['SmallH2']))
        elif clean.startswith('- '):
            story.append(Paragraph('&#8226; ' + clean[2:], styles['TinyBody']))
        else:
            story.append(Paragraph(clean, styles['TinyBody']))
    doc.build(story)
    return str(out)


def main():
    items = load_items()
    master = ROOT / 'GSN_IN_REAL_LIFE_MASTER_CAPABILITY_BANK.md'
    write_md(master, TITLE + ': Master Capability Bank', 'A complete modular bank of the real-life capabilities GSN can explain. It stretches beyond the old 23 items because the current product story has more practical doors than one public paper should carry.', items)

    generated = [master]
    for key, (label, intro) in AUDIENCES.items():
        selected = [item for item in items if key in item['aud']]
        path = ROOT / f'{COPY_NAME[key]}.md'
        write_md(path, f'{TITLE}: {label}', intro, selected, selected=True)
        generated.append(path)

    with (ROOT / 'GSN_IN_REAL_LIFE_SELECTION_MAP.csv').open('w', newline='', encoding='utf-8') as fh:
        writer = csv.writer(fh)
        writer.writerow(['id','title','audiences','status_boundary'])
        for item in items:
            writer.writerow([item['id'], item['title'], ', '.join(item['aud']), item['status']])
    (ROOT / 'GSN_IN_REAL_LIFE_SELECTION_MAP.json').write_text(json.dumps(items, indent=2), encoding='utf-8')

    readme = ROOT / 'README.md'
    readme.write_text(f"# {TITLE}\n\nThis folder contains the real-life capability bank and audience-specific copies. The master has {len(items)} modules. The sector copies are intentionally smaller so each audience receives only what concerns them.\n\nUse the verified self-service guide as the truth source for live status, routes and setup steps. This folder is for practical explanation, sales conversations and pilot demos.\n\nFiles generated:\n" + ''.join(f"- {p.name}\n" for p in generated) + "- GSN_IN_REAL_LIFE_SELECTION_MAP.csv\n- GSN_IN_REAL_LIFE_SELECTION_MAP.json\n", encoding='utf-8')

    rendered = []
    for md in generated:
        rendered.append(simple_docx(md))
        rendered.append(simple_pdf(md))
    print(json.dumps({'modules': len(items), 'markdown_files': [str(p) for p in generated], 'rendered': rendered}, indent=2))

if __name__ == '__main__':
    main()
